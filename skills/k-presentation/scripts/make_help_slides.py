"""Prepend hidden "cheat sheet" slides to a template.

The pattern is borrowed from the official École polytechnique template, which
ships three instruction slides — palette, master usage, text levels — at the
top of the file, carrying `show="0"`: present in the document, absent from the
running order. Some authors deliberately keep them at the top of every deck
they build, so the rules travel with the file.

This script generalises that to any template: it reads the theme palette, the
master's body levels and the full layout inventory, and writes them onto hidden
slides in the template's own colours. You then always have, inside the file,
the answer to "which layout, which placeholder idx, which colour".

It is idempotent: slides it generated are tagged in their notes and replaced on
the next run, so re-running after a template revision is safe.

Usage
-----
    uv run --with python-pptx python make_help_slides.py Template_DxO.pptx
    uv run --with python-pptx python make_help_slides.py deck.pptx --visible
    uv run --with python-pptx python make_help_slides.py deck.pptx -o out.pptx
"""

from __future__ import annotations

import argparse
import pathlib
import re
import sys
import zipfile

MARKER = "k-presentation:help"

#: Printed on the grammar sheet. Kept short on purpose: the sheet is a
#: reminder inside the file, not the documentation.
GRAMMAR = (
    "# front-matter (real YAML)",
    "---",
    "layout: Double bloc de texte   # exact layout name",
    "part: III \u00b7 The mission       # breadcrumb, optional",
    "title: What this slide claims",
    "figure: value_loop             # or image: ../path.png",
    "caption: source, method note",
    "quote: a pull quote",
    "---",
    "",
    "# body",
    "- A main idea                 level 0",
    "  - a qualifier               level 1 (2 spaces)",
    "- date | label                cells, split on ' | '",
    "",
    "## Named block                a column",
    "## Notes                      speaker notes, never shown",
    "",
    "# inline emphasis",
    "**strong**  __accent__  *stress*",
    "",
    "# hard rules",
    "two bullet levels max \u2014 a third means two ideas",
    "nothing projected below 16 pt",
    "numbers in filenames are sort keys, not page numbers",
)

#: The composition patterns worth repeating in almost any deck.
ARCHETYPES = (
    ("Cover", "title, author, entity, dates \u2014 host logo most visible"),
    ("Sommaire", "level-1 titles only, no page references"),
    ("Statement", "one idea, large type, nothing else"),
    ("Two columns", "a comparison, a rule at the gutter"),
    ("Figure + takeaway", "a drawn diagram, one sentence under it"),
    ("Image + one line", "full-bleed, treated to the house palette"),
    ("Object + satellites", "a centred visual spine, text blocks around it"),
    ("Table", "a real table, readable during questions"),
    ("Timeline", "date | label per entry"),
    ("Conclusion", "synthesis plus an opening; stays on screen"),
)
EMU = 914400
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


# ------------------------------------------------------------- reading ---
def read_theme(path: str) -> tuple[dict, list[str]]:
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if re.match(r"ppt/theme/theme\d+\.xml", n)]
        if not names:
            return {}, []
        xml = z.read(min(names)).decode("utf-8", "replace")
    colours: dict[str, str] = {}
    scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.DOTALL)
    if scheme:
        pattern = (
            r'<a:(\w+)>\s*<a:(?:srgbClr val="([0-9A-Fa-f]{6})"'
            r'|sysClr[^>]*lastClr="([0-9A-Fa-f]{6})")'
        )
        for m in re.finditer(pattern, scheme.group(0)):
            colours[m.group(1)] = (m.group(2) or m.group(3)).upper()
    fonts: list[str] = []
    fs = re.search(r"<a:fontScheme.*?</a:fontScheme>", xml, re.DOTALL)
    if fs:
        for face in re.findall(r'typeface="([^"]+)"', fs.group(0)):
            if face and face not in fonts:
                fonts.append(face)
    return colours, fonts[:4]


def read_levels(master) -> list[tuple[str, str, str, str]]:
    from pptx.oxml.ns import qn

    styles = master._element.find(qn("p:txStyles"))
    if styles is None:
        return []
    body = styles.find(qn("p:bodyStyle"))
    if body is None:
        return []
    rows = []
    for i, level in enumerate(body, start=1):
        rpr = level.find(qn("a:defRPr"))
        latin = rpr.find(qn("a:latin")) if rpr is not None else None
        size = rpr.get("sz") if rpr is not None else None
        bold = " bold" if rpr is not None and rpr.get("b") == "1" else ""
        bu = level.find(qn("a:buChar"))
        rows.append(
            (
                f"Level {i}",
                f"{int(size) / 100:g} pt{bold}" if size else "—",
                (latin.get("typeface") if latin is not None else "inherited"),
                f"marL={level.get('marL') or 0}  bullet="
                f"{bu.get('char') if bu is not None else 'none'}",
            )
        )
    return rows[:6]


# ------------------------------------------------------------- drawing ---
def blank_layout(prs):
    """The emptiest layout available: help slides must not inherit furniture."""
    layouts = list(prs.slide_masters[0].slide_layouts)
    for layout in layouts:
        if re.search(r"\b(blank|vide)\b", layout.name, re.IGNORECASE):
            return layout
    return min(layouts, key=lambda x: len(list(x.placeholders)))


class Sheet:
    """A help slide being drawn. Plain text boxes only — no placeholders, so
    the sheet never depends on the layout it sits on."""

    def __init__(self, prs, ink: str, accent: str, font: str, mono: str):
        from pptx.util import Inches

        self.prs = prs
        self.slide = prs.slides.add_slide(blank_layout(prs))
        self.w = prs.slide_width / EMU
        self.h = prs.slide_height / EMU
        self.ink, self.accent, self.font, self.mono = ink, accent, font, mono
        self.y = 0.42
        self._Inches = Inches

    def _box(self, x, y, w, h):
        return (self._Inches(x), self._Inches(y), self._Inches(w), self._Inches(h))

    def fit(self, count: int, *, bottom: float = 0.55) -> float:
        """Line pitch that makes `count` lines fit below the cursor.

        Slide heights differ between templates — 7.5 in for one of the bundled
        pair, 5.625 for the other — so a fixed pitch overflows the shorter one
        and the last lines fall off the page. Derive it instead.
        """
        room = max(0.4, self.h - self.y - bottom)
        return min(0.26, room / max(1, count))

    def text(
        self,
        s,
        *,
        size=11,
        x=0.5,
        w=None,
        colour=None,
        bold=False,
        font=None,
        dy=None,
        y=None,
    ):
        from pptx.dml.color import RGBColor
        from pptx.util import Pt

        w = w if w is not None else self.w - 1.0
        top = self.y if y is None else y
        shape = self.slide.shapes.add_textbox(
            *self._box(x, top, w, max(0.22, size / 72 * 1.5))
        )
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = s
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font or self.font
        run.font.color.rgb = RGBColor.from_string(colour or self.ink)
        if y is None:
            self.y += dy if dy is not None else size / 72 * 1.55
        return shape

    def swatch(self, x, y, name, hexval, *, size=0.62):
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE

        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, *self._box(x, y, size, size)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(hexval)
        shape.line.color.rgb = RGBColor.from_string("BFBFBF")
        shape.shadow.inherit = False
        shape.text_frame.text = ""
        self.text(name, size=8, x=x, w=size + 0.5, y=y + size + 0.04)
        self.text(
            f"#{hexval}", size=7, x=x, w=size + 0.5, colour="7F7F7F", y=y + size + 0.20
        )

    def rule(self, *, dy=0.16):
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_CONNECTOR
        from pptx.util import Inches, Pt

        line = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(0.5),
            Inches(self.y),
            Inches(self.w - 0.5),
            Inches(self.y),
        )
        line.line.color.rgb = RGBColor.from_string(self.accent)
        line.line.width = Pt(1.0)
        self.y += dy

    def finish(self, hidden: bool):
        self.slide.notes_slide.notes_text_frame.text = (
            f"{MARKER}\nGenerated by .agents/skills/k-presentation/scripts/"
            f"make_help_slides.py — re-run it to refresh. Delete nothing by "
            f"hand: the script replaces its own slides."
        )
        if hidden:
            self.slide._element.set("show", "0")
        return self.slide


# ---------------------------------------------------------------- build ---
def purge_previous(prs) -> int:
    """Remove slides this script generated on an earlier run."""
    lst = prs.slides._sldIdLst
    doomed = []
    for sld_id, slide in zip(list(lst), list(prs.slides)):
        if slide.has_notes_slide and MARKER in slide.notes_slide.notes_text_frame.text:
            doomed.append(sld_id)
    for sld_id in doomed:
        prs.part.drop_rel(sld_id.attrib.get(REL))
        lst.remove(sld_id)
    return len(doomed)


def move_to_front(prs, count: int) -> None:
    """python-pptx appends; help sheets belong at the top.

    They are inserted AFTER any leading run of already-hidden slides, so a
    template that ships official instruction slides keeps them in first
    position. Ours come next, then the content.
    """
    lst = prs.slides._sldIdLst
    ids = list(lst)
    fresh = ids[-count:]
    kept = ids[:-count]
    offset = 0
    for sld_id, slide in zip(kept, list(prs.slides)[: len(kept)]):
        if slide._element.get("show") == "0":
            offset += 1
        else:
            break
    for sld_id in fresh:
        lst.remove(sld_id)
    for i, sld_id in enumerate(fresh):
        lst.insert(offset + i, sld_id)


def build(path: str, out: str | None, hidden: bool) -> pathlib.Path:
    from pptx import Presentation

    prs = Presentation(path)
    colours, fonts = read_theme(path)
    ink = colours.get("dk1", "1B1F24")
    accent = colours.get("accent1", "808080")
    font = fonts[0] if fonts else "Arial"
    mono = "Consolas"
    master = prs.slide_masters[0]

    removed = purge_previous(prs)

    # --- sheet 1 : palette, fonts, text levels ---------------------------
    s = Sheet(prs, ink, accent, font, mono)
    s.text(
        "Template cheat sheet — palette & type",
        size=20,
        bold=True,
        colour=accent,
        dy=0.42,
    )
    s.text(
        "Hidden slide. It travels with the file so the rules never get "
        "lost. Regenerate with make_help_slides.py.",
        size=9,
        colour="7F7F7F",
        dy=0.34,
    )
    s.rule()
    s.text(
        "Theme colours — the only legitimate source of a brand colour",
        size=11,
        bold=True,
        dy=0.30,
    )
    if colours:
        shown = list(colours.items())[:12]
        row_y = s.y
        per_row = 6
        for i, (name, hexval) in enumerate(shown):
            s.swatch(
                0.5 + (i % per_row) * 1.30, row_y + (i // per_row) * 1.15, name, hexval
            )
        rows = (len(shown) + per_row - 1) // per_row
        s.y = row_y + rows * 1.15 + 0.12
    s.rule()
    s.text(
        f"Fonts — {' · '.join(fonts) if fonts else 'not declared'}",
        size=11,
        bold=True,
        dy=0.30,
    )
    for name, size, face, extra in read_levels(master):
        s.text(
            f"{name:<9}  {size:<14}  {face:<18}  {extra}", size=9, font=mono, dy=0.22
        )
    s.text(
        "If every marL is 0, this template carries hierarchy by SIZE "
        "ONLY: set marL/indent yourself and add a bullet glyph, or "
        "indentation is invisible on screen.",
        size=9,
        colour=accent,
        dy=0.30,
    )
    s.finish(hidden)
    sheets = 1

    # --- sheets 2+ : layout inventory ------------------------------------
    layouts = list(master.slide_layouts)
    # Two text lines per layout plus the header block: derive how many fit.
    usable = prs.slide_height / EMU - 1.9
    per_sheet = max(4, int(usable / 0.46))
    for start in range(0, len(layouts), per_sheet):
        chunk = layouts[start : start + per_sheet]
        s = Sheet(prs, ink, accent, font, mono)
        n = start // per_sheet + 1
        total = (len(layouts) + per_sheet - 1) // per_sheet
        s.text(
            f"Layout inventory ({n}/{total}) — name, then placeholder idx",
            size=20,
            bold=True,
            colour=accent,
            dy=0.42,
        )
        s.text(
            "Use the exact name in a slide's `layout:` field. Placeholder "
            "idx values are NOT in reading order — pair them by "
            "coordinates (inspect_template.py prints those).",
            size=9,
            colour="7F7F7F",
            dy=0.34,
        )
        s.rule()
        pitch = s.fit(2 * len(chunk))
        for layout in chunk:
            idxs = ", ".join(
                f"{p.placeholder_format.idx}"
                f":{str(p.placeholder_format.type).split(' ')[0]}"
                for p in sorted(
                    layout.placeholders, key=lambda q: q.placeholder_format.idx
                )
            )
            s.text(
                layout.name, size=min(11, pitch * 72 * 1.0), bold=True, dy=pitch * 0.82
            )
            s.text(
                idxs or "(no placeholder — draw everything yourself)",
                size=min(8, pitch * 72 * 0.72),
                font=mono,
                colour="595959",
                dy=pitch * 1.18,
            )
        s.finish(hidden)
        sheets += 1

    # --- last sheet : the grammar and the archetypes on offer -------------
    s = Sheet(prs, ink, accent, font, mono)
    s.text(
        "How to drive this template — grammar & archetypes",
        size=20,
        bold=True,
        colour=accent,
        dy=0.42,
    )
    s.text(
        "One slide = one markdown file. Python places, never writes. "
        "Full contract: .agents/skills/k-presentation/SKILL.md",
        size=9,
        colour="7F7F7F",
        dy=0.34,
    )
    s.rule()
    left = 0.5
    right = s.w / 2 + 0.15
    top = s.y
    pitch = s.fit(len(GRAMMAR))
    size = min(8.5, pitch * 72 * 0.62)
    for line in GRAMMAR:
        s.text(
            line,
            size=size,
            x=left,
            w=s.w / 2 - 0.7,
            font=mono if not line.startswith("#") else font,
            colour=ink if not line.startswith("#") else accent,
            bold=line.startswith("#"),
            dy=pitch,
        )
    s.y = top
    pitch = s.fit(2 * len(ARCHETYPES) + 3)
    s.text(
        "Archetypes worth repeating",
        size=min(11, pitch * 72 * 1.3),
        bold=True,
        x=right,
        w=s.w / 2 - 0.7,
        dy=pitch * 1.6,
    )
    for name, why in ARCHETYPES:
        s.text(
            f"·  {name}",
            size=min(9.5, pitch * 72 * 1.05),
            x=right,
            w=s.w / 2 - 0.7,
            bold=True,
            dy=pitch * 0.82,
        )
        s.text(
            f"   {why}",
            size=min(8, pitch * 72 * 0.88),
            x=right,
            w=s.w / 2 - 0.7,
            colour="595959",
            dy=pitch * 1.18,
        )
    s.finish(hidden)
    sheets += 1

    move_to_front(prs, sheets)
    target = pathlib.Path(out) if out else pathlib.Path(path)
    # Save to a sibling then replace: writing a zip over the very file
    # python-pptx still holds open makes zipfile emit "Duplicate name"
    # warnings and can leave stale parts behind.
    tmp = target.with_suffix(target.suffix + ".tmp")
    prs.save(tmp)
    tmp.replace(target)
    print(
        f"{target.name}: {sheets} help slide(s) "
        f"{'hidden' if hidden else 'visible'} at the top"
        f"{f' (replaced {removed})' if removed else ''}, "
        f"{len(prs.slides._sldIdLst)} slide(s) total, "
        f"{target.stat().st_size / 1e6:.2f} MB"
    )
    return target


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("path")
    ap.add_argument("-o", "--out", help="write elsewhere (default: in place)")
    ap.add_argument(
        "--visible", action="store_true", help="do not mark the sheets hidden"
    )
    args = ap.parse_args(argv)
    try:
        import pptx  # noqa: F401
    except ImportError:
        print(
            "python-pptx missing. Try:\n  uv run --with python-pptx python "
            "make_help_slides.py …",
            file=sys.stderr,
        )
        return 2
    build(args.path, args.out, hidden=not args.visible)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
